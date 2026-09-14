"""Build editable manuscript and expand computed tables from frozen CSV files."""
from pathlib import Path
import csv,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
O=Path('output');E=list(csv.DictReader((O/'ruptures.csv').open()));S=json.loads((O/'summary.json').read_text())
def mdtable(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(map(str,r))+' |' for r in rows])
gloss=[('Prime gap','gₙ=pₙ₊₁−pₙ.'),('Step','sₙ=gₙ/2 for consecutive odd primes.'),('Prime-step tree','Centred block representation of odd primes, with seed 1.'),('Record step','A step larger than every earlier step.'),('Record expansion','A transition establishing a new record step.'),('Sequential record expansion','Record increase of exactly one step level.'),('Prime-gap rupture','Record increase of more than one step level.'),('Rupture jump','Jₙ=sₙ−Mₙ.'),('Skipped step','An intermediate level bypassed by a rupture.'),('Skipped-step set','Kₙ={Mₙ+1,…,sₙ−1}.'),('Recovery','First later occurrence of a skipped step.'),('Complete recovery','First upper endpoint by which every skipped step has appeared.'),('Recovery delay','A stated index, numerical, logarithmic or event-count distance.'),('Right-censored recovery','Not observed through the complete recovery-data boundary.'),('Rupture desert','The prime-free integer interior of a rupture interval.'),('Prime-factor oasis','Descriptive view of composite factor structure in that interior.'),('Computational boundary','Largest declared point of complete relevant data.')]
T={
'GLOSSARY':'Table 1. Short glossary.\n\n'+mdtable(['Term','Meaning'],gloss),
'EARLY':'Table 2. Early record steps.\n\n'+mdtable(['Pair','Gap','Step','Classification'],[['3→5',2,1,'Initial record'],['7→11',4,2,'Sequential'],['23→29',6,3,'Sequential'],['89→97',8,4,'Sequential'],['113→127',14,7,'Rupture']]),
'FIRST':'Table 3. First-rupture calculation.\n\n'+mdtable(['Quantity','Value'],[['Pair','113→127'],['Gap','127−113=14'],['Step','14/2=7'],['Earlier maximum','4'],['Condition','7>4+1'],['Jump','7−4=3'],['Skipped set','{5,6}'],['Skipped count','3−1=2']]),
'BOUNDARY':'Table 4. Detection and recovery boundaries. Both endpoints must be within each bound.\n\n'+mdtable(['Purpose','Boundary','Coverage'],[['Original rupture detection','5,000,000','13 events; independently reproduced'],['Direct detection and recovery','100,000,000','16 events; full local prime enumeration'],['Published detection and recovery','10²⁰','68 events; source-reported exhaustive coverage'],['Candidate 69','Above 10²⁰','Local pair verification only']]),
'RUNTIME':f"The final numerical run used Python {S['python']} on {S['platform']} ({S['cpu']}). It took {S['seconds']:.2f} seconds and approximately {S['peak_MiB']:.1f} MiB peak resident memory. This is a single-process CPU computation, excluding downloads, spreadsheet extraction, figures and document rendering. No specialised hardware or GPU is required; wall time varies by machine.",
'ORIGINAL':'Table 5. Original thirteen events, independently reproduced. M is the earlier maximum.\n\n'+mdtable(['No.','Lower p','Upper q','s','M','Skipped levels'],[[e['event'],e['p'],e['q'],e['step'],e['previous'],e['skipped_start']+('–'+e['skipped_end'] if e['skipped_start']!=e['skipped_end'] else '')] for e in E[:13]]),
'RECOVERY13':'Table 6. Complete recovery of the first thirteen events. Delay counts subsequent ruptures.\n\n'+mdtable(['No.','Last recovery pair','Last step','Delay'],[[e['event'],e['complete_p']+'→'+e['complete_q'],e['last_step'],e['delay_ruptures']] for e in E[:13]]),
'CENSORING':'Table 7. Recovery and censoring status at 10²⁰.\n\n'+mdtable(['Events','Recovered levels','Unresolved','Status'],[['1–59','All','0','Complete'],['60–66','All','0','Complete'],['67','22 of 51','29','Right-censored'],['68','1 of 23','22','Right-censored'],['69','Not assessed','Not applicable','Outside detection boundary']]),
'CLASSIFICATION':'Table 8. Classification of conclusions.\n\n'+mdtable(['Class','Conclusion'],[['Proved statements','Every rupture is a record; skipped count equals jump minus one; skipped sets are disjoint.'],['Verified data','Local sieve and pair checks; published-priority catalogue through 10²⁰.'],['Finite observations','66 complete recoveries; observed maximum delay eight.'],['Descriptive interpretations','Prime-step tree, rupture desert and prime-factor oasis.'],['Potential utility','Classification, cataloguing, model comparison and teaching.'],['Conjectures','No new conjecture proposed; universal recovery is unproved.'],['Unresolved questions','51 skipped levels censored; candidate 69 priority; factor-richness.']]),
'CATALOGUE':'Table A1. Expanded catalogue. C=confirmed within 10²⁰; P=outside-boundary candidate.\n\n'+mdtable(['No.','Lower prime p','Gap','M','s','Skipped','Status'],[[e['event'],e['p'],e['gap'],e['previous'],e['step'],e['skipped_start']+('–'+e['skipped_end'] if e['skipped_start']!=e['skipped_end'] else ''),'C' if e['status']=='confirmed' else 'P'] for e in E]),
'PROVENANCE':'Table A2. Sources and provenance. All external datasets accessed 13 September 2026.\n\n'+mdtable(['Source','Role','Limit'],[['Original PDF and workbook','Origin and audit targets','Not assumed correct'],['Prime Gap List Project [1–2]','Record and first-occurrence source','Exhaustive priority adopted only through 10²⁰'],['OEIS A002386/A000101 [3]','85 record endpoint cross-checks','Shared underlying discoveries'],['OEIS A000230 [4]','First-occurrence endpoints','Steps 1–721 cross-checked'],['OEIS A014321 [5]','Half-gap chronological order','All 747 available terms cross-checked'],['Independent segmented sieve','Full enumeration through 10⁸','Not extended to 10²⁰'],['Bounded Miller–Rabin [7]','All listed record and accepted recovery pairs','Consecutiveness only']])}
src=O/'Prime_Gap_Ruptures.md';text=Path('manuscript_template.md').read_text()
for k,v in T.items():text=text.replace('{{'+k+'}}',v)
src.write_text(text)
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
normal=d.styles['Normal'];normal.font.name='Cambria';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.05
for name,size in [('Title',22),('Heading 1',14),('Heading 2',12)]:
 st=d.styles[name];st.font.name='Cambria';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
for st in d.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.core_properties.author='Riccardo Panza';d.core_properties.title=text.splitlines()[0][2:]
d.core_properties.subject='An observational and computational study of prime-gap ruptures'
d.core_properties.keywords='prime gaps, record gaps, maximal gaps, first occurrences, rupture, recovery, computational number theory'
d.core_properties.comments="Release 1.0.0. Prepared with disclosed AI assistance under the author's direction."
foot=sec.footer.paragraphs[0];foot.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
lines=text.splitlines();i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   cells=[c.strip() for c in lines[i].strip('|').split('|')]
   if not all(re.fullmatch(r'[-:]+',c) for c in cells):rows.append(cells)
   i+=1
  chunks=[rows] if len(rows)<=51 else [[rows[0]]+rows[1:51],[rows[0]]+rows[51:]]
  for chunk_no,chunk in enumerate(chunks):
   if chunk_no:d.add_page_break()
   table=d.add_table(rows=1,cols=len(chunk[0]));table.style='Table Grid';table.autofit=False
   if len(chunk[0])==7: widths=[.35,2.45,.5,.5,.5,1.05,.65]
   elif len(chunk[0])==6:widths=[.45,1.4,1.4,.6,.6,1.7]
   elif len(chunk[0])==2:widths=[1.8,5.3]
   else:widths=[7.1/len(chunk[0])]*len(chunk[0])
   for c,w in zip(table.columns,widths):c.width=Inches(w)
   for j,row in enumerate(chunk):
    cells=table.rows[0].cells if j==0 else table.add_row().cells
    for cell,value,w in zip(cells,row,widths):
     cell.width=Inches(w);cell.text=value
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
      if len(chunk)<=18 and j<len(chunk)-1:p.paragraph_format.keep_with_next=True
      for r in p.runs:r.font.size=Pt(9 if len(row)>5 else 10);r.bold=j==0
     if j==0:
      shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6E6E6');cell._tc.get_or_add_tcPr().append(shade)
    trpr=table.rows[j]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
    if j==0:trpr.append(OxmlElement('w:tblHeader'))
   d.add_paragraph()
  continue
 if line.startswith('# '):d.add_paragraph(line[2:],style='Title')
 elif line.startswith('## '):d.add_heading(line[3:],level=1)
 elif line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line);p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(O/m[2]),width=Inches(5.4 if '08_first13' in m[2] else 6.8));c=d.add_paragraph(m[1]);c.paragraph_format.space_after=Pt(10)
  for r in c.runs:r.italic=True;r.font.size=Pt(9)
 else:
  p=d.add_paragraph(line)
  if line.startswith('Table '):p.paragraph_format.keep_with_next=True
 i+=1
d.save(O/'Prime_Gap_Ruptures.docx')
print('Saved manuscript',len(d.paragraphs),'paragraphs',len(d.tables),'tables')
